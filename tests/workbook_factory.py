"""Builds a small fictional workbook that reproduces the real file's quirks:

- Data Base: header in row 1, a visit spanning 3 rows (merged cells + follow-up comment rows).
- Active Sheet: two blank rows before the header, an inserted row that shifts IDs, a
  duplicated ID, a record without an ID, US-style text dates and an Excel serial date.
- To share with Dealer: a subset, with one conflicting status.
- Dealer shop sheets with a duplicate, and a pivot sheet that must be skipped.
All names and numbers are made up.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

BASE = [
    "ID", "Start time", "Completion time", "Email", "Name", "Last modified time",
    "Date of Site Visit", "Name of Salesperson", "Name of Area (Location) of Site", "City",
    "Name of Potential Site", "Address", "Name of Owners", "Contractor", "Architect",
    "Contact Person Name", "Contact Details", "Name of Company",
    "Estimated Paint Start Time Line", "Google Location", "Comments (If Any)",
]  # fmt: skip
DB_HEADERS = [*BASE, "SAR", "Dealer To Route", "Jotun Comments"]
ACTIVE_HEADERS = [
    *BASE, "SAR", "Dealer To Route", "Jotun Comments", "Site Visited by Dealer/ODS",
    "Quotation Shared", "Won/\nLost", "First Order received", "Supplying Dealer Name",
    "Quote Value", "Status",
]  # fmt: skip
SHARE_HEADERS = [
    *BASE, "Dealer To Route", "Quotation Shared", "Won/\nLost", "Supplying Dealer Name",
    "Quote Value", "Status",
]  # fmt: skip


def site(
    sid, name, phone, *, visit=datetime(2025, 3, 1), start=datetime(2025, 3, 2, 10, 0),
    salesperson="Jane Wanjiru", area="Lavington", city="Nairobi ", owner="Owner " + "A",
    contractor="Acme Builders Ltd", comment="Walls being plastered.", estimate=datetime(2025, 5, 1),
    **extra,
):  # fmt: skip
    row = {
        "ID": sid, "Start time": start, "Completion time": start, "Email": "anonymous",
        "Date of Site Visit": visit, "Name of Salesperson": salesperson,
        "Name of Area (Location) of Site": area, "City": city, "Name of Potential Site": name,
        "Address": f"{area} road", "Name of Owners": owner, "Contractor": contractor,
        "Architect": "TBC", "Contact Person Name": "Mr. Contact", "Contact Details": phone,
        "Name of Company": "N/A", "Estimated Paint Start Time Line": estimate,
        "Google Location": "https://maps.app.goo.gl/example", "Comments (If Any)": comment,
    }  # fmt: skip
    row.update(extra)
    return row


def _write(ws, headers, rows, start_row):
    for col, name in enumerate(headers, start=1):
        ws.cell(start_row, col, name)
    r = start_row
    for row in rows:
        r += 1
        if row is None:  # continuation row: only a follow-up comment
            ws.cell(r, headers.index("Comments (If Any)") + 1, "Follow-up: samples delivered.")
            continue
        for col, name in enumerate(headers, start=1):
            key = "Won/ Lost" if name == "Won/\nLost" else name
            if key in row:
                ws.cell(r, col, row[key])
    return r


def build_workbook(path: Path) -> Path:
    wb = Workbook()
    db = wb.active
    db.title = "Data Base"
    db_rows = [
        site(
            1, "Kamau proposed villa", "0712345678", SAR="Antony", **{"Dealer To Route": "Cidone"}
        ),
        site(2, "Sunrise Apartments", 722111222, salesperson="Jane Wanjiru - project."),
        site(3, "Baraka Heights", "0733 444555 or 0711222333", comment="Roofing done."),
        None,  # rows 5-6 continue ID 3 (merged cells in the real file)
        None,
        site(4, "Green Court", "TBA", city="Edoret", visit=datetime(2024, 8, 7)),
        site(5, "Lakeview Towers", "+254700999888", contractor="acme builders ltd"),
    ]
    last = _write(db, DB_HEADERS, db_rows, start_row=1)
    db.merge_cells(start_row=4, start_column=1, end_row=6, end_column=1)
    assert last == 8

    active = wb.create_sheet("Active Sheet")
    active_rows = [
        site(
            1, "Kamau proposed villa", "0712345678", **{"Won/ Lost": "won", "Status": "Following"}
        ),
        site(2, "Sunrise Apartments", "0722111222", start="3/16/2025 11:07", visit="3/16/2025"),
        # Inserted row: takes ID 3, so the real ID 3 below is a duplicate.
        site(3, "Ocean Breeze Hotel", "0799000111", start=45716.5380208333),
        site(3, "Baraka Heights", "0733444555", **{"Quote Value": 450000}),
        site(None, "Pine Ridge Villa", None, salesperson="Peter Otieno", estimate="15.10.25"),
        site(5, "Lakeview Towers", "0700999888"),
    ]
    _write(active, ACTIVE_HEADERS, active_rows, start_row=3)

    share = wb.create_sheet("To share with Dealer")
    share_rows = [
        site(1, "Kamau proposed villa", "0712345678", **{"Status": "Not Following"}),
        site(5, "Lakeview Towers", "0700999888"),
    ]
    _write(share, SHARE_HEADERS, share_rows, start_row=3)

    for title, header in (("Philmax data", "ID"), ("Sheet3", "Sr. No.")):
        ws = wb.create_sheet(title)
        ws.append([header, "Brand Ambasador Shop", "General Location", "Exact Location"])
        ws.append([1, "Shop One", "Kiambu", "Ruaka"])
        ws.append([2, "Shop Two", "Nairobi", "Kilimani"])
        if title == "Philmax data":
            ws.append([3, "shop one", "Kiambu", "Ruaka"])  # duplicate of row 1

    pivot = wb.create_sheet("Hit Rate")
    pivot.append(["Row Labels", "Lost", "Won"])
    wb.save(path)
    return path

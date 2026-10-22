import json
from types import SimpleNamespace
from unittest.mock import AsyncMock

from openpyxl import load_workbook
from sqlalchemy import select

from brains_bot.bot import brochure_handlers, outreach_handlers, partner_handlers
from brains_bot.database.models import Brochure, CatalogProduct, Contact
from brains_bot.database.session import session_scope
from brains_bot.services import brochures, catalog
from brains_bot.services.analysis import AnalysisError
from brains_bot.services.categories import sync_categories
from tests.conftest import ALLOWED_ID, make_update
from tests.test_chat import FakeTransport, text

LINES = [
    "SAUTI SOUND LTD - Outdoor and Home Audio. Call 0733 444 555",
    "1. Garden Rock Speakers (pair) - weatherproof - KES 68,000",
    "2. Pool-side Bollard Speaker - from KES 42,500 each",
    "3. Multi-room Streaming Amplifier - price on request",
]
PRODUCTS = {"supplier": "Sauti Sound Ltd", "title": "Outdoor audio",
            "categories": ["outdoor_audio", "audio", "kitchens"],
            "products": [
    {"category": "outdoor_audio", "name": "Garden Rock Speakers (pair)", "brand": None,
     "description": "weatherproof", "price": "KES 68,000", "price_unit": "pair"},
    {"category": "outdoor_audio", "name": "Pool-side Bollard Speaker", "brand": None,
     "description": "IP66", "price": "from KES 42,500", "price_unit": "each"},
    {"category": "audio", "name": "Multi-room Streaming Amplifier", "brand": None,
     "description": "4 zones", "price": None, "price_unit": None},
    {"category": "kitchens", "name": "Not a category", "price": None},
]}  # fmt: skip


def text_pdf(lines):
    """A minimal one-page PDF with a text layer."""
    stream = (
        "BT /F1 11 Tf 40 800 Td 14 TL "
        + " ".join(f"({line.replace('(', '[').replace(')', ']')}) '" for line in lines)
        + " ET"
    )
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = "%PDF-1.4\n", []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n"
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF"
    return out.encode("latin-1")


class MistralTransport(FakeTransport):
    """The OCR endpoint is rate limited (as on smaller plans); chat calls are played back."""

    provider = SimpleNamespace(key="mistral")

    def _send(self, method, path, json=None):
        if path == "/ocr":
            self.requests.append(json)
            raise AnalysisError("Mistral is rate limited right now.")
        return super()._send(method, path, json)


def test_text_pdfs_are_read_locally_and_photos_by_the_ai():
    transport = MistralTransport()
    reading = brochures.read_document(transport, text_pdf(LINES), "application/pdf")
    assert reading.method == "text" and "KES 68,000" in reading.text
    assert transport.requests == []  # nothing left the bot

    transport = MistralTransport(text("Garden Rock Speakers KES 68,000"))
    reading = brochures.read_document(transport, b"\xff\xd8 photo", "image/jpeg")
    assert reading.method == "vision" and "68,000" in reading.text
    assert transport.requests[0]["model"] == "mistral-ocr-latest"  # tried first

    try:
        brochures.read_document(FakeTransport(), b"photo", "image/jpeg")
    except AnalysisError as exc:
        assert "MISTRAL_API_KEY" in str(exc)
    else:
        raise AssertionError("DeepSeek cannot read pictures")


def test_products_become_a_sheet_to_check_and_load(factory, tmp_path):
    with session_scope(factory) as s:
        sync_categories(s)
        s.add(CatalogProduct(sku="P-1", category_key="audio", name="Prime pair",
                             supplier="Prime Audio", install_stage="any time"))  # fmt: skip
        s.add(CatalogProduct(sku="S-OLD", category_key="audio", name="Old Sauti amp",
                             supplier="Sauti Sound Ltd", install_stage="any time"))  # fmt: skip
        names = dict(s.execute(select(CatalogProduct.sku, CatalogProduct.name)).all())
    transport = FakeTransport(text(json.dumps(PRODUCTS)))
    info = brochures.extract_products(transport, "\n".join(LINES), {"audio": "Audio",
                                      "outdoor_audio": "Outdoor Audio"}, "Sauti Sound")  # fmt: skip
    sent = transport.requests[0]["messages"][1]["content"]
    assert "0733" not in sent and "Sauti Sound" in sent
    assert info.categories == ["outdoor_audio", "audio"] and len(info.products) == 3
    with session_scope(factory) as s:
        path = brochures.catalog_sheet(s, info, "Sauti Sound Ltd", "sauti.pdf", tmp_path / "c.xlsx")
    rows = list(load_workbook(path).active.iter_rows(values_only=True))
    assert [r[1] for r in rows[2:]] == ["S-OLD", "SAUTIS-001", "SAUTIS-002", "SAUTIS-003"]
    assert rows[3][6] == "68000" and rows[4][7] == "each (from)" and rows[5][6] is None
    assert rows[3][8] == "finishing" and "check price" in rows[3][12]
    with session_scope(factory) as s:
        stats = catalog.load_catalog(s, path)
        assert stats.products == 4 and stats.deactivated == 0
        active = set(s.scalars(select(CatalogProduct.sku).where(CatalogProduct.active)))
        assert active == {"P-1", "S-OLD", "SAUTIS-001", "SAUTIS-002", "SAUTIS-003"}
    assert names["P-1"] == "Prime pair"


def _file_update(caption=""):
    update = make_update(ALLOWED_ID)
    update.effective_message.photo = []
    update.effective_message.video = None
    update.effective_message.caption = caption
    update.effective_message.document = SimpleNamespace(
        file_id="FILE1", file_unique_id="U1", file_name="sauti.pdf",
        mime_type="application/pdf", file_size=5000,
    )  # fmt: skip
    return update


def _tap(data):
    query = SimpleNamespace(data=data, answer=AsyncMock(), edit_message_text=AsyncMock())
    tap = make_update(ALLOWED_ID)
    tap.callback_query = query
    tap.effective_chat = SimpleNamespace(send_document=AsyncMock(), send_photo=AsyncMock())
    return tap, query


def _context(context, factory, *answers):
    context.bot_data["db"] = factory
    context.bot_data["drafter"] = SimpleNamespace(transport=MistralTransport(*answers))
    context.user_data = {}
    pdf = text_pdf(LINES)
    context.bot = SimpleNamespace(get_file=AsyncMock(return_value=SimpleNamespace(
        download_as_bytearray=AsyncMock(return_value=bytearray(pdf)))))  # fmt: skip


def _buttons(update):
    markup = update.effective_message.reply_text.await_args.kwargs["reply_markup"]
    return {b.callback_data.split(":")[1]: b.callback_data for row in markup.inline_keyboard
            for b in row}  # fmt: skip


async def test_catalog_button_returns_a_sheet_and_files_the_brochure(factory, context):
    with session_scope(factory) as s:
        sync_categories(s)
    _context(context, factory, text(json.dumps(PRODUCTS)))
    update = _file_update("Sauti Sound")
    await brochure_handlers.file_received(update, context)
    assert "What is this file?" in update.effective_message.reply_text.await_args.args[0]
    tap, query = _tap(_buttons(update)["catalog"])
    await brochure_handlers.file_callback(tap, context)
    reply = query.edit_message_text.await_args.args[0]
    assert "Read 3 products (2 with a price)" in reply and "Outdoor Audio" in reply
    sheet = tap.effective_chat.send_document.await_args
    assert sheet.kwargs["filename"] == "catalog_Sauti_Sound.xlsx"
    assert "send this file back to me to load it" in sheet.kwargs["caption"]
    with session_scope(factory) as s:
        (saved,) = s.scalars(select(Brochure)).all()
        assert saved.supplier == "Sauti Sound" and saved.file_id == "FILE1"
        assert saved.category_keys == ["outdoor_audio", "audio"]

    # Find it again and get the file to forward.
    update = make_update(ALLOWED_ID)
    context.args = ["outdoor", "audio"]
    await brochure_handlers.brochures_command(update, context)
    (send,) = _buttons(update).values()
    assert "Outdoor audio · Sauti Sound (PDF)" in str(
        update.effective_message.reply_text.await_args.kwargs["reply_markup"].inline_keyboard
    )
    tap, _ = _tap(send)
    await brochure_handlers.send_callback(tap, context)
    tap.effective_chat.send_document.assert_awaited_once()
    assert tap.effective_chat.send_document.await_args.args[0] == "FILE1"

    # A draft for that category offers the brochure to forward with the message.
    draft = SimpleNamespace(id=7, text="Hi", sent_at=None, replied_at=None)
    project = SimpleNamespace(id=3, name="Villa")
    category = SimpleNamespace(key="outdoor_audio", name="Outdoor Audio")
    _, markup = outreach_handlers.draft_message(project, category, None, draft, 1)
    assert markup.inline_keyboard[-1][0].callback_data == "brochure:cat:outdoor_audio"
    tap, _ = _tap("brochure:cat:outdoor_audio")
    await brochure_handlers.send_callback(tap, context)
    assert tap.effective_chat.send_document.await_args.args[0] == "FILE1"


async def test_partner_button_adds_the_partner_with_the_brochure(factory, context):
    with session_scope(factory) as s:
        sync_categories(s)
    partner = {"partners": [{"name": None, "company": "Sauti Sound Ltd", "area": "Nairobi",
                             "trade": "supplier", "categories": ["outdoor_audio"],
                             "rates": []}]}  # fmt: skip
    _context(context, factory, text(json.dumps(partner)))
    update = _file_update()
    await brochure_handlers.file_received(update, context)
    tap, query = _tap(_buttons(update)["partner"])
    await brochure_handlers.file_callback(tap, context)
    preview = query.edit_message_text.await_args.args[0]
    assert "PARTNER: Sauti Sound Ltd · Nairobi" in preview and "Work: Outdoor Audio" in preview
    save = query.edit_message_text.await_args.kwargs["reply_markup"].inline_keyboard[0][0]
    tap, query = _tap(save.callback_data)
    await partner_handlers.partner_callback(tap, context)
    with session_scope(factory) as s:
        contact = s.scalar(select(Contact).where(Contact.partner.is_(True)))
        brochure = s.scalar(select(Brochure))
        assert brochure.contact_id == contact.id and brochure.category_keys == ["outdoor_audio"]
        assert contact.phone_e164 == "+254733444555"


def test_other_files_are_not_brochures():
    update = _file_update()
    update.effective_message.document.mime_type = "application/zip"
    update.effective_message.document.file_name = "x.zip"
    assert not brochure_handlers.is_brochure(update)


def screen_recording(colours, seconds_each=1.5, fps=10):
    """A small MP4: each colour held still for a while, like pausing on a catalogue screen."""
    import io

    import av
    from PIL import Image

    out = io.BytesIO()
    with av.open(out, "w", format="mp4") as container:
        stream = container.add_stream("libx264", rate=fps)
        stream.width, stream.height, stream.pix_fmt = 180, 320, "yuv420p"
        for colour in colours:
            image = Image.new("RGB", (180, 320), colour)
            for _ in range(int(seconds_each * fps)):
                container.mux(stream.encode(av.VideoFrame.from_image(image)))
        container.mux(stream.encode())
    return out.getvalue()


def test_a_screen_recording_is_read_frame_by_frame_once_per_screen():
    video = screen_recording(["white", "black", "white"])
    frames = brochures.video_frames(video)
    assert len(frames) == 3  # one per screen, not one per second
    assert all(f.startswith(b"\xff\xd8") for f in frames)  # JPEG

    transport = MistralTransport(*(text(f"Screen {i} KES 1,000") for i in range(3)))
    reading = brochures.read_video(transport, video)
    assert reading.method == "vision" and reading.pages == 3 and "Screen 2" in reading.text
    ocr_tries = [r for r in transport.requests if r.get("model") == "mistral-ocr-latest"]
    assert len(ocr_tries) == 1  # not tried again for every frame

    try:
        brochures.video_frames(b"not a video")
    except AnalysisError as exc:
        assert "could not be opened" in str(exc)
    else:
        raise AssertionError("a broken video is reported")


async def test_a_catalogue_video_becomes_a_sheet(factory, context, monkeypatch):
    with session_scope(factory) as s:
        sync_categories(s)
    _context(context, factory, text(json.dumps(PRODUCTS)))
    monkeypatch.setattr(
        brochures,
        "read_video",
        lambda transport, data: brochures.Reading("\n".join(LINES), 5, "vision"),
    )
    update = _file_update("Sauti Sound")
    update.effective_message.document = None
    update.effective_message.video = SimpleNamespace(
        file_id="VID1", file_unique_id="UV1", file_name=None, file_size=4_000_000
    )
    await brochure_handlers.file_received(update, context)
    assert "frame by frame" in update.effective_message.reply_text.await_args.args[0]
    tap, query = _tap(_buttons(update)["catalog"])
    await brochure_handlers.file_callback(tap, context)
    prompt = context.bot_data["drafter"].transport.requests[0]["messages"][1]["content"]
    assert "WhatsApp catalogue" in prompt  # the AI is told frames overlap
    sheet = tap.effective_chat.send_document.await_args
    assert "products from the catalogue recording" in sheet.kwargs["caption"]
    with session_scope(factory) as s:
        (saved,) = s.scalars(select(Brochure)).all()
        assert saved.kind == "video" and brochures.label(saved).endswith("(catalogue video)")

    tap, _ = _tap(f"brochure:one:{saved.id}")
    tap.effective_chat.send_video = AsyncMock()
    await brochure_handlers.send_callback(tap, context)
    assert tap.effective_chat.send_video.await_args.args[0] == "VID1"


def test_repeated_products_are_listed_once():
    answer = dict(PRODUCTS, products=PRODUCTS["products"] * 2)
    info = brochures.extract_products(
        FakeTransport(text(json.dumps(answer))), "text", {"outdoor_audio": "Outdoor Audio",
                                                          "audio": "Audio"}
    )  # fmt: skip
    assert len(info.products) == len({p.name for p in info.products})

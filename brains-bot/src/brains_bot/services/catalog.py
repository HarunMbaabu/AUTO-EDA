"""The verified product catalog: load it from a spreadsheet and match products to a project.

The catalog file has a 'Catalog' sheet with one row per product (see the template in the
README). A file is checked completely before anything is stored: one bad row rejects the whole
file, so a half-loaded catalog never reaches the sales team. Each supplier's products are
replaced by the suppliers named in the new file: a file with only Prime Audio rows updates Prime
Audio and leaves other suppliers' products as they are.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from brains_bot.database.models import CatalogProduct, Category
from brains_bot.importers.normalize import redact_contact_details

SHEET = "Catalog"
REQUIRED = ("category", "sku", "name", "install_stage")
COLUMNS = (
    "category", "sku", "name", "brand", "supplier", "description", "price_kes", "price_unit",
    "install_stage", "retrofit_ok", "available", "price_checked", "notes",
)  # fmt: skip

# The latest construction stage (services/stages.py) at which a product can still be installed
# the normal way. Later than that, only products marked retrofit_ok fit.
INSTALL_STAGES = {"pre-wiring": 4, "first fix": 4, "finishing": 6, "any time": 7}
_STAGE_ALIASES = {"prewiring": "pre-wiring", "first_fix": "first fix", "firstfix": "first fix",
                  "anytime": "any time", "any_time": "any time"}  # fmt: skip
_MULTIPLIERS = {"k": 1_000, "m": 1_000_000}


class CatalogError(ValueError):
    """The file cannot be loaded; the message lists what to fix."""


@dataclass
class CatalogStats:
    products: int = 0
    by_category: dict[str, int] = field(default_factory=dict)
    deactivated: int = 0
    unpriced: int = 0
    suppliers: list[str] = field(default_factory=list)


def supplier_key(name: str | None) -> str:
    return re.sub(r"[^a-z0-9]+", " ", (name or "").lower()).strip()


def is_catalog_file(path: Path) -> bool:
    wb = load_workbook(path, read_only=True)
    try:
        return SHEET in wb.sheetnames
    finally:
        wb.close()


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_price(value: Any) -> Decimal | None:
    """300000, '300,000', 'KES 300K', '1.5M' -> Decimal; blank -> None."""
    if value is None or _text(value) == "":
        return None
    if isinstance(value, int | float | Decimal):
        return Decimal(str(value))
    text = _text(value).lower().replace("kes", "").replace("ksh", "").replace(",", "")
    text = text.replace("from", "").strip()
    match = re.fullmatch(r"(\d+(?:\.\d+)?)\s*([km]?)", text)
    if not match:
        raise ValueError(f"price '{_text(value)}' is not a number")
    try:
        number = Decimal(match.group(1))
    except InvalidOperation as exc:
        raise ValueError(f"price '{_text(value)}' is not a number") from exc
    return number * _MULTIPLIERS.get(match.group(2), 1)


def _install_stage(value: Any) -> str:
    text = _text(value).lower().replace("–", "-")
    text = _STAGE_ALIASES.get(text, text)
    if text not in INSTALL_STAGES:
        raise ValueError(
            f"install_stage '{_text(value)}' must be one of: {', '.join(INSTALL_STAGES)}"
        )
    return text


def _yes_no(value: Any) -> bool:
    text = _text(value).lower()
    if text in ("yes", "y", "true", "1"):
        return True
    if text in ("", "no", "n", "false", "0"):
        return False
    raise ValueError(f"retrofit_ok '{_text(value)}' must be yes or no")


def _date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if _text(value) == "":
        return None
    try:
        return date.fromisoformat(_text(value)[:10])
    except ValueError as exc:
        raise ValueError(f"price_checked '{_text(value)}' must be a date (YYYY-MM-DD)") from exc


def read_catalog(path: Path, category_keys: set[str]) -> list[dict[str, Any]]:
    """Every product row, validated. Raises CatalogError listing the problems."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET not in wb.sheetnames:
            raise CatalogError(f"the file has no '{SHEET}' sheet")
        rows = list(wb[SHEET].iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        raise CatalogError(f"the '{SHEET}' sheet is empty")
    header = [_text(h).lower() for h in rows[0]]
    missing = [c for c in REQUIRED if c not in header]
    if missing:
        raise CatalogError(f"missing column(s): {', '.join(missing)}")
    index = {name: header.index(name) for name in COLUMNS if name in header}

    products, problems, seen = [], [], set()
    for number, row in enumerate(rows[1:], start=2):
        cells = {name: row[i] if i < len(row) else None for name, i in index.items()}
        if all(_text(v) == "" for v in cells.values()):
            continue
        if _text(cells["category"]).lower().startswith("category key"):
            continue  # the template's explanation row
        try:
            category = _text(cells["category"]).lower()
            if category not in category_keys:
                raise ValueError(f"unknown category '{_text(cells['category'])}'")
            sku = _text(cells["sku"])
            if not sku:
                raise ValueError("sku is empty")
            if sku.lower() in seen:
                raise ValueError(f"sku '{sku}' appears twice")
            seen.add(sku.lower())
            name = _text(cells["name"])
            if not name:
                raise ValueError("name is empty")
            products.append(
                {
                    "sku": sku,
                    "category_key": category,
                    "name": name,
                    "brand": _text(cells.get("brand")) or None,
                    "supplier": _text(cells.get("supplier")) or None,
                    "description": _text(cells.get("description")),
                    "price_kes": parse_price(cells.get("price_kes")),
                    "price_unit": _text(cells.get("price_unit")) or None,
                    "install_stage": _install_stage(cells["install_stage"]),
                    "retrofit_ok": _yes_no(cells.get("retrofit_ok")),
                    "available": _text(cells.get("available")) or None,
                    "price_checked": _date(cells.get("price_checked")),
                    "notes": _text(cells.get("notes")),
                }
            )
        except ValueError as exc:
            problems.append(f"row {number}: {exc}")
    if problems:
        shown = "\n".join(f"- {p}" for p in problems[:10])
        more = f"\n…and {len(problems) - 10} more" if len(problems) > 10 else ""
        raise CatalogError(f"{len(problems)} row(s) need fixing:\n{shown}{more}")
    if not products:
        raise CatalogError("no products found")
    return products


def load_catalog(session: Session, path: Path) -> CatalogStats:
    categories = {c.key: c for c in session.scalars(select(Category).where(Category.active))}
    rows = read_catalog(path, set(categories))
    existing = {p.sku.lower(): p for p in session.scalars(select(CatalogProduct))}
    stats = CatalogStats()
    for row in rows:
        product = existing.pop(row["sku"].lower(), None)
        if product is None:
            product = CatalogProduct(sku=row["sku"])
            session.add(product)
        for key, value in row.items():
            setattr(product, key, value)
        product.active, product.source_file = True, path.name
        stats.by_category[row["category_key"]] = stats.by_category.get(row["category_key"], 0) + 1
        stats.unpriced += row["price_kes"] is None
    suppliers = {supplier_key(row.get("supplier")) for row in rows}
    stats.suppliers = sorted({row.get("supplier") or "(no supplier)" for row in rows})
    for product in existing.values():
        if product.active and supplier_key(product.supplier) in suppliers:
            product.active = False
            stats.deactivated += 1
    stats.products = len(rows)
    session.flush()
    stocked = set(
        session.scalars(select(CatalogProduct.category_key).where(CatalogProduct.active).distinct())
    )
    for key, category in categories.items():
        category.has_verified_catalog = key in stocked
    session.flush()
    return stats


def products_for(session: Session, category_key: str) -> list[CatalogProduct]:
    return list(
        session.scalars(
            select(CatalogProduct)
            .where(CatalogProduct.active, CatalogProduct.category_key == category_key)
            .order_by(CatalogProduct.price_kes.asc().nulls_last(), CatalogProduct.id)
        )
    )


def fits_stage(product: CatalogProduct, stage: int | None) -> bool:
    """Can this product still go in at the given construction stage? Unknown stage: yes."""
    return stage is None or stage <= INSTALL_STAGES[product.install_stage] or product.retrofit_ok


def needs_retrofit(product: CatalogProduct, stage: int | None) -> bool:
    """True when the project is past the product's normal install stage."""
    return stage is not None and stage > INSTALL_STAGES[product.install_stage]


def suggestions(session: Session, category_key: str, stage: int | None, limit: int = 5):
    """Catalog products that still fit the project's stage, those installed the normal way
    first, then retrofit options, cheapest first within each."""
    fitting = [p for p in products_for(session, category_key) if fits_stage(p, stage)]
    fitting.sort(key=lambda p: needs_retrofit(p, stage))
    return fitting[:limit]


def price_text(product: CatalogProduct) -> str:
    if product.price_kes is None:
        return "price to be scoped"
    unit = product.price_unit or ""
    prefix = "from " if "from" in unit.lower() else ""
    unit = re.sub(r"\s*\(?from\)?", "", unit, flags=re.I).strip()
    return f"{prefix}KES {product.price_kes:,.0f}" + (f" {unit}" if unit else "")


def product_summary(product: CatalogProduct, stage: int | None = None) -> dict[str, Any]:
    """What the AI and chat see about a product: catalog facts only."""
    summary = {
        "sku": product.sku,
        "name": product.name,
        "brand": product.brand,
        "supplier": product.supplier,
        "description": product.description,
        "price": price_text(product),
        "install_stage": product.install_stage,
        "retrofit_ok": product.retrofit_ok,
        "available": product.available,
        "price_checked": product.price_checked.isoformat() if product.price_checked else None,
    }
    # Catalog rows are free text too: a supplier's phone or email must not reach the AI.
    for key, value in summary.items():
        if isinstance(value, str):
            summary[key] = redact_contact_details(value)
    if stage is not None:
        summary["fits_current_stage_as"] = (
            "retrofit" if needs_retrofit(product, stage) else "normal install"
        )
    return summary

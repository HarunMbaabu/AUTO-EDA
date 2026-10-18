"""Read worksheets into logical records, keeping sheet/row provenance and original values."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

PROJECT_SHEETS = ("Data Base", "Active Sheet", "To share with Dealer")
DEALER_SHEETS = ("Philmax data", "Sheet3")
SUMMARY_SHEETS = ("Sheet1", "Hit Rate", "Pivot (2)")

# A row with none of these filled (and no ID) continues the record above it: in "Data Base"
# one visit can span several rows, with the extra rows holding follow-up comments only.
KEY_COLUMNS = ("Name of Potential Site", "Name of Salesperson", "Date of Site Visit")


@dataclass
class SourceRow:
    sheet: str
    first_row: int
    last_row: int
    values: dict[str, Any]
    extra: list[tuple[str, Any]] = field(default_factory=list)  # (column, value) from later rows
    issues: list[str] = field(default_factory=list)

    @property
    def source_id(self) -> int | None:
        value = self.values.get("ID")
        if isinstance(value, int | float) and not isinstance(value, bool):
            return int(value)
        if isinstance(value, str) and value.strip().isdigit():
            return int(value.strip())
        return None

    def get(self, column: str) -> Any:
        return self.values.get(column)

    def raw(self) -> dict[str, Any]:
        """JSON-safe copy of the original values, for the source_records table."""

        def safe(v: Any) -> Any:
            if isinstance(v, datetime | date):
                return v.isoformat()
            return v

        data = {k: safe(v) for k, v in self.values.items() if v is not None}
        if self.extra:
            data["_continuation"] = [[k, safe(v)] for k, v in self.extra]
        return data


def _header_name(value: Any) -> str | None:
    if value is None:
        return None
    return re.sub(r"\s+", " ", str(value)).strip() or None


def find_header_row(ws: Worksheet, max_scan: int = 10) -> int:
    for row in range(1, min(ws.max_row, max_scan) + 1):
        names = [_header_name(c.value) for c in ws[row]]
        if names and names[0] in ("ID", "Sr. No.") and sum(1 for n in names if n) >= 3:
            return row
    raise ValueError(f"no header row found in sheet {ws.title!r}")


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def read_sheet(ws: Worksheet) -> list[SourceRow]:
    header_row = find_header_row(ws)
    headers = [_header_name(c.value) for c in ws[header_row]]
    records: list[SourceRow] = []
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
    ):
        values = {h: v for h, v in zip(headers, row, strict=False) if h}
        if all(_is_blank(v) for v in values.values()):
            continue
        has_key = any(not _is_blank(values.get(k)) for k in KEY_COLUMNS)
        is_project_sheet = ws.title in PROJECT_SHEETS
        if is_project_sheet and _is_blank(values.get("ID")) and not has_key and records:
            parent = records[-1]
            parent.extra.extend((k, v) for k, v in values.items() if not _is_blank(v))
            parent.last_row = row_idx
            continue
        record = SourceRow(sheet=ws.title, first_row=row_idx, last_row=row_idx, values=values)
        if record.source_id is None and is_project_sheet:
            record.issues.append("missing source ID")
        records.append(record)
    return records


@dataclass
class Workbook:
    sheets: dict[str, list[SourceRow]]
    skipped_sheets: list[str]
    unknown_sheets: list[str]


def read_workbook(path: Path) -> Workbook:
    wb = load_workbook(path, data_only=True, read_only=False)
    sheets: dict[str, list[SourceRow]] = {}
    skipped, unknown = [], []
    for ws in wb.worksheets:
        if ws.title in PROJECT_SHEETS or ws.title in DEALER_SHEETS:
            sheets[ws.title] = read_sheet(ws)
        elif ws.title in SUMMARY_SHEETS:
            skipped.append(ws.title)
        else:
            unknown.append(ws.title)
    missing = [s for s in PROJECT_SHEETS[:1] if s not in sheets]
    if missing:
        raise ValueError(f"workbook is missing required sheet(s): {missing}")
    return Workbook(sheets=sheets, skipped_sheets=skipped, unknown_sheets=unknown)
